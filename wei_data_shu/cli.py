"""Command-line interface for wei-data-shu."""

from __future__ import annotations

import argparse
import csv
import io
import json
import os
import sys
import unicodedata
from collections.abc import Mapping
from html import escape
from pathlib import Path
from typing import Any, Sequence

from wei_data_shu.utils import generate_password, search_colors
from wei_data_shu.utils.textio import bom_tolerant_encoding, read_text

_DATABASE_HINT = "缺少 database 依赖，请先安装: pip install 'wei-data-shu[database]'"
_DB_PASSWORD_ENV = "WEI_DATA_SHU_DB_PASSWORD"
_MAIL_PASSWORD_ENV = "WEI_DATA_SHU_MAIL_PASSWORD"


def _ensure_utf8_stdout() -> None:
    reconfigure = getattr(sys.stdout, "reconfigure", None)
    if callable(reconfigure):
        reconfigure(encoding="utf-8")


def _emit(text: str) -> None:
    """输出文本并保证以换行结尾，避免与 shell 提示符粘连。"""
    print(text, end="" if text.endswith("\n") else "\n")


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(prog="wei-data-shu", description="Utilities for wei-data-shu")
    subparsers = parser.add_subparsers(dest="command")

    password_parser = subparsers.add_parser("password", help="Generate readable passwords")
    password_parser.add_argument("-l", "--length", type=int, default=13, help="Password length")
    password_parser.add_argument("-c", "--count", type=int, default=1, help="Number of passwords to generate")

    colors_parser = subparsers.add_parser("colors", help="View or search color palette")
    colors_parser.add_argument("query", nargs="?", help="Search by hex, English name, or Chinese name")

    date_parser = subparsers.add_parser("date", help="Print a date offset by N days from today")
    date_parser.add_argument("-d", "--days", type=int, default=0, help="Days to subtract from today (default 0)")
    date_parser.add_argument("-f", "--format", default="%Y-%m-%d", help="strftime format (default %%Y-%%m-%%d)")

    excel_parser = subparsers.add_parser("excel", help="Inspect an Excel workbook")
    excel_sub = excel_parser.add_subparsers(dest="excel_command", required=True)
    info_parser = excel_sub.add_parser("info", help="List sheets and row counts of a workbook")
    info_parser.add_argument("file", help="Path to the .xlsx file")

    convert_parser = subparsers.add_parser("convert", help="Convert a document (source: md / docx / pptx / xlsx)")
    convert_parser.add_argument("source", help="Source document (.md / .markdown / .docx / .pptx / .xlsx / .xlsm)")
    convert_parser.add_argument("-o", "--output", required=True, help="Target document path (.md / .docx / .pptx)")
    convert_parser.add_argument("--no-overwrite", action="store_true", help="Fail when the target already exists")

    md_parser = subparsers.add_parser("md", help="Print a document as Markdown")
    md_parser.add_argument("file", help="Document path (.md / .markdown / .docx / .pptx / .xlsx / .xlsm)")
    md_parser.add_argument("-o", "--output", help="Write Markdown to a file instead of stdout")

    files_parser = subparsers.add_parser("files", help="Locate the most recently created folder")
    files_sub = files_parser.add_subparsers(dest="files_command", required=True)
    latest_parser = files_sub.add_parser("latest", help="Print the most recently created subfolder")
    latest_parser.add_argument("directory", help="Directory whose subfolders are compared")

    md2html_parser = subparsers.add_parser("md2html", help="Render a document as HTML")
    md2html_parser.add_argument("file", help="Document path (.md / .markdown / .docx / .pptx / .xlsx / .xlsm)")
    md2html_parser.add_argument("-o", "--output", help="Write HTML to a file instead of stdout")
    md2html_parser.add_argument(
        "--fragment", action="store_true", help="Emit a bare HTML fragment instead of a full page"
    )
    md2html_parser.add_argument("--title", action="store_true", help="Also render the document title as a heading")

    text_parser = subparsers.add_parser("text", help="Clean and renumber plain text")
    text_sub = text_parser.add_subparsers(dest="text_command", required=True)
    clean_parser = text_sub.add_parser("clean", help="De-duplicate and renumber numbered lines")
    clean_parser.add_argument("file", help="Path to a plain-text file")
    clean_parser.add_argument(
        "-m",
        "--mode",
        choices=["renumber", "original", "sql"],
        default="renumber",
        help=(
            "renumber: 重新编号并去重，同内容不同编号也算重复（默认）；"
            "original: 保留原编号，仅整行完全相同才去重；"
            "sql: 转为 SQL IN 列表"
        ),
    )
    clean_parser.add_argument("-o", "--output", help="Write the result to a file instead of stdout")
    clean_parser.add_argument("--encoding", default="utf-8", help="Input encoding (default utf-8)")

    table_parser = subparsers.add_parser("table", help="Convert between Markdown tables and CSV/TSV")
    table_parser.add_argument("file", help="Source file (.md / .markdown / .csv / .tsv)")
    table_parser.add_argument("-o", "--output", help="Write the result to a file instead of stdout")
    table_parser.add_argument("-i", "--index", type=int, default=0, help="Which Markdown table to take (default 0)")
    table_parser.add_argument("--delimiter", help="CSV delimiter (default: ',' for .csv, tab for .tsv)")
    table_parser.add_argument("--encoding", default="utf-8", help="File encoding (default utf-8)")

    data_parser = subparsers.add_parser("data", help="Inspect a data file")
    data_sub = data_parser.add_subparsers(dest="data_command", required=True)
    data_info_parser = data_sub.add_parser("info", help="Print shape, per-column null counts and the first rows")
    data_info_parser.add_argument("file", help="Data file (.csv / .tsv / .txt / .json / .xlsx)")
    data_info_parser.add_argument("-n", "--head", type=int, default=5, help="Rows to preview, 0 to skip (default 5)")

    plot_parser = subparsers.add_parser("plot", help="Render a chart to an image file")
    plot_parser.add_argument("file", help="Data file (.csv / .tsv / .txt / .json / .xlsx)")
    plot_parser.add_argument("-o", "--output", required=True, help="Output image path (for example chart.png)")
    plot_parser.add_argument(
        "-k",
        "--kind",
        choices=["line", "bar", "hist", "box", "scatter", "pie", "heatmap"],
        default="line",
        help="Chart type (default line)",
    )
    plot_parser.add_argument("-x", "--x", help="Column for the x axis / grouping / category")
    plot_parser.add_argument("-y", "--y", help="Column(s) for the y axis / values; comma separated for several")
    plot_parser.add_argument("-t", "--title", help="Chart title")
    plot_parser.add_argument("--horizontal", action="store_true", help="Draw bar charts horizontally")

    mail_parser = subparsers.add_parser("mail", help="Send email through SMTP/SSL")
    mail_sub = mail_parser.add_subparsers(dest="mail_command", required=True)
    mail_send_parser = mail_sub.add_parser("send", help="Send a plain-text or HTML email")
    mail_send_parser.add_argument("--host", required=True, help="SMTP host")
    mail_send_parser.add_argument("--port", type=int, default=465, help="SMTP SSL port (default 465)")
    mail_send_parser.add_argument("--user", required=True, help="Sender account, used as the From address")
    mail_send_parser.add_argument("--password", help=f"Sender password (or set {_MAIL_PASSWORD_ENV})")
    mail_send_parser.add_argument("-t", "--to", action="append", required=True, help="Recipient (repeatable)")
    mail_send_parser.add_argument("-s", "--subject", required=True, help="Subject line")
    mail_send_parser.add_argument("-b", "--body", help="Inline body text")
    mail_send_parser.add_argument("--body-file", help="Read the body from a file")
    mail_send_parser.add_argument("--html", action="store_true", help="Send the body as HTML")
    mail_send_parser.add_argument("-a", "--attach", action="append", help="Attachment path (repeatable)")
    mail_send_parser.add_argument("--dry-run", action="store_true", help="Compose and validate without sending")
    mail_send_parser.add_argument("--encoding", default="utf-8", help="Encoding for --body-file (default utf-8)")

    db_parser = subparsers.add_parser("db", help="Run a MySQL query")
    db_sub = db_parser.add_subparsers(dest="db_command", required=True)
    db_query_parser = db_sub.add_parser("query", help="Run one query and print the rows")
    db_query_parser.add_argument("--sql", required=True, help="SQL statement to run")
    db_query_parser.add_argument("--host", default="127.0.0.1", help="MySQL host (default 127.0.0.1)")
    db_query_parser.add_argument("--port", type=int, default=3306, help="MySQL port (default 3306)")
    db_query_parser.add_argument("--user", required=True, help="MySQL user")
    db_query_parser.add_argument("--password", help=f"MySQL password (or set {_DB_PASSWORD_ENV})")
    db_query_parser.add_argument("--database", required=True, help="Database name")
    db_query_parser.add_argument(
        "-f",
        "--format",
        choices=["table", "csv", "json"],
        default="table",
        help="Output format (default table)",
    )

    return parser


def _run_password(args: argparse.Namespace) -> int:
    if args.count <= 0:
        raise ValueError("count must be greater than 0")
    for _ in range(args.count):
        print(generate_password(args.length))
    return 0


def _run_colors(args: argparse.Namespace) -> int:
    results = search_colors(args.query)
    if not results:
        print(f'No colors matched "{args.query}".')
        return 1

    for record in results:
        index = record["index"]
        color_hex = record["hex"]
        name = record["name"]
        name_zh = record["name_zh"]
        print(f"{index:>2}. {color_hex} | {name} | {name_zh}")
    return 0


def _run_date(args: argparse.Namespace) -> int:
    from wei_data_shu.text import DateFormat

    print(DateFormat(interval_day=args.days, timeclass="date").get_timeparameter(Format=args.format))
    return 0


def _run_excel(args: argparse.Namespace) -> int:
    if args.excel_command != "info":
        return 1
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(args.file, read_only=True, data_only=True)
        try:
            for sheet_name in workbook.sheetnames:
                print(f"{sheet_name}\t{workbook[sheet_name].max_row} 行")
        finally:
            workbook.close()
    except ImportError:
        print("缺少 excel 依赖，请先安装: pip install 'wei-data-shu[excel]'")
        return 1
    return 0


def _run_convert(args: argparse.Namespace) -> int:
    try:
        from wei_data_shu.docs import convert

        target = convert(args.source, args.output, overwrite=not args.no_overwrite)
    except (ImportError, ValueError, FileNotFoundError, FileExistsError) as exc:
        print(f"转换失败: {exc}")
        return 1
    print(f"已生成: {target}")
    return 0


def _run_md(args: argparse.Namespace) -> int:
    try:
        from wei_data_shu.docs import read_doc

        text = read_doc(args.file).to_markdown()
    except (ImportError, ValueError, FileNotFoundError) as exc:
        print(f"读取失败: {exc}")
        return 1

    if args.output:
        Path(args.output).write_text(text, encoding="utf-8")
        print(f"已生成: {args.output}")
    else:
        _emit(text)
    return 0


def _run_files(args: argparse.Namespace) -> int:
    if args.files_command != "latest":
        return 1

    from wei_data_shu.files import FileManagement

    base = Path(args.directory)
    if not base.is_dir():
        print(f"目录不存在: {base}")
        return 1

    latest = FileManagement().find_latest_folder(base)
    if latest is None:
        print(f"{base} 下没有子目录。")
        return 1

    print(latest)
    return 0


def _html_page(body: str, title: str | None) -> str:
    """把 HTML 片段包成可直接用浏览器打开的完整页面。"""
    heading = escape(title or "", quote=False)
    return (
        '<!DOCTYPE html>\n<html lang="zh-CN">\n<head>\n<meta charset="utf-8">\n'
        f"<title>{heading}</title>\n</head>\n<body>\n{body}\n</body>\n</html>\n"
    )


def _run_md2html(args: argparse.Namespace) -> int:
    try:
        from wei_data_shu.docs import read_doc, render_html

        document = read_doc(args.file)
    except ImportError as exc:
        print(f"缺少依赖: {exc}")
        return 1
    except (ValueError, FileNotFoundError, OSError) as exc:
        print(f"读取失败: {exc}")
        return 1

    fragment = render_html(document, include_title=args.title)
    text = fragment if args.fragment else _html_page(fragment, document.title)

    if args.output:
        Path(args.output).write_text(text, encoding="utf-8")
        print(f"已生成: {args.output}")
    else:
        _emit(text)
    return 0


def _run_text(args: argparse.Namespace) -> int:
    if args.text_command != "clean":
        return 1

    from wei_data_shu.text import StringBaba, textCombing

    try:
        text = read_text(args.file, args.encoding)
    except OSError as exc:
        print(f"读取失败: {exc}")
        return 1

    if args.mode == "sql":
        result = StringBaba(text).format_string_sql()
    else:
        # "重排" 先剥离编号再比较，因此同内容不同编号也算重复，随后统一重编号；
        # "原版" 保留原编号，只有整行文本完全相同时才去重。
        combiner = textCombing(global_var1="原版" if args.mode == "original" else "重排")
        result = combiner.format_text(text)

    if args.output:
        Path(args.output).write_text(result + "\n", encoding="utf-8")
        print(f"已生成: {args.output}")
    else:
        _emit(result)
    return 0


def _run_table(args: argparse.Namespace) -> int:
    from wei_data_shu.docs import Table, parse_markdown, render_table

    source = Path(args.file)
    if not source.is_file():
        print(f"文件不存在: {source}")
        return 1

    suffix = source.suffix.lower()
    if suffix in {".md", ".markdown"}:
        try:
            text = read_text(source, args.encoding)
        except OSError as exc:
            print(f"读取失败: {exc}")
            return 1

        tables = [node for node in parse_markdown(text).nodes if isinstance(node, Table)]
        if not tables:
            print("文档中未找到 Markdown 表格。")
            return 1
        if args.index >= len(tables):
            print(f"--index={args.index} 越界，文档中共有 {len(tables)} 个表格。")
            return 1

        buffer = io.StringIO()
        csv.writer(buffer, lineterminator="\n").writerows(tables[args.index].normalized())
        result = buffer.getvalue()
    elif suffix in {".csv", ".tsv"}:
        delimiter = args.delimiter or ("\t" if suffix == ".tsv" else ",")
        try:
            with source.open("r", encoding=bom_tolerant_encoding(args.encoding), newline="") as handle:
                rows = [row for row in csv.reader(handle, delimiter=delimiter)]
        except OSError as exc:
            print(f"读取失败: {exc}")
            return 1
        result = render_table(rows) + "\n"
    else:
        print(f"不支持的文件类型: {suffix or '(无扩展名)'}。支持: .md / .markdown / .csv / .tsv")
        return 1

    if args.output:
        Path(args.output).write_text(result, encoding="utf-8")
        print(f"已生成: {args.output}")
    else:
        _emit(result)
    return 0


def _split_columns(value: str | None) -> list[str] | None:
    if not value:
        return None
    return [item.strip() for item in value.split(",") if item.strip()]


def _first_column(value: str | None) -> str | None:
    columns = _split_columns(value)
    return columns[0] if columns else None


def _read_frame(path: str) -> Any:
    """读取数据文件，缺依赖时抛 :class:`ImportError`。"""
    from wei_data_shu.analysis._deps import require_deps

    require_deps("pandas")
    from wei_data_shu.analysis import read_any

    return read_any(path)


def _run_data(args: argparse.Namespace) -> int:
    if args.data_command != "info":
        return 1

    try:
        frame = _read_frame(args.file)
    except ImportError as exc:
        print(f"依赖不可用: {exc}")
        return 1
    except (ValueError, FileNotFoundError, OSError) as exc:
        print(f"读取失败: {exc}")
        return 1

    row_count = len(frame)
    print(f"文件: {args.file}")
    print(f"行数: {row_count}  列数: {len(frame.columns)}")

    if len(frame.columns):
        print()
        print("列\t非空\t缺失\t类型")
        for name in frame.columns:
            series = frame[name]
            missing = int(series.isna().sum())
            print(f"{name}\t{row_count - missing}\t{missing}\t{series.dtype}")

    if args.head > 0 and row_count:
        print()
        print(f"前 {min(args.head, row_count)} 行:")
        print(frame.head(args.head).to_string())
    return 0


def _run_plot(args: argparse.Namespace) -> int:
    try:
        from wei_data_shu.analysis._deps import require_deps

        require_deps("pandas", "matplotlib")
        from wei_data_shu.analysis import (
            plot_bar,
            plot_box,
            plot_corr_heatmap,
            plot_hist,
            plot_line,
            plot_pie,
            plot_scatter,
        )

        frame = _read_frame(args.file)
    except ImportError as exc:
        print(f"依赖不可用: {exc}")
        return 1
    except (ValueError, FileNotFoundError, OSError) as exc:
        print(f"读取失败: {exc}")
        return 1

    kind = args.kind
    if kind == "bar" and not args.x:
        print("bar 图需要 --x 指定分组列。")
        return 1
    if kind == "scatter" and not (args.x and args.y):
        print("scatter 图需要 --x 与 --y 指定列。")
        return 1
    if kind == "pie" and not args.x:
        print("pie 图需要 --x 指定分类列。")
        return 1
    if kind == "line" and args.y and not args.x:
        print("line 图使用 --y 时需同时指定 --x，否则会绘制所有列。")
        return 1

    columns = _split_columns(args.y)
    target = args.output
    try:
        if kind == "line":
            plot_line(frame, x=args.x, y=columns, title=args.title, save_path=target)
        elif kind == "bar":
            plot_bar(frame, x=args.x, y=columns, title=args.title, horizontal=args.horizontal, save_path=target)
        elif kind == "hist":
            plot_hist(frame, col=columns, title=args.title, save_path=target)
        elif kind == "box":
            plot_box(frame, cols=columns, title=args.title, save_path=target)
        elif kind == "scatter":
            plot_scatter(frame, x=args.x, y=_first_column(args.y), title=args.title, save_path=target)
        elif kind == "pie":
            plot_pie(frame, col=_first_column(args.x), title=args.title, save_path=target)
        else:
            plot_corr_heatmap(frame, cols=columns, title=args.title, save_path=target)
    except ImportError as exc:
        print(f"依赖不可用: {exc}")
        return 1
    except (ValueError, KeyError) as exc:
        print(f"绘图失败: {exc}")
        return 1

    print(f"已生成: {target}")
    return 0


def _run_mail(args: argparse.Namespace) -> int:
    if args.mail_command != "send":
        return 1

    password = args.password or os.environ.get(_MAIL_PASSWORD_ENV)
    if not password:
        print(f"缺少密码: 用 --password 传入，或设置环境变量 {_MAIL_PASSWORD_ENV}。")
        return 1

    body = args.body
    if args.body_file:
        try:
            body = read_text(args.body_file, args.encoding)
        except OSError as exc:
            print(f"读取正文失败: {exc}")
            return 1
    if body is None:
        print("缺少正文: 用 --body 或 --body-file 指定。")
        return 1

    from wei_data_shu.mail import DailyEmailReport, MailError

    attachments = [Path(item) for item in args.attach or []]
    report = DailyEmailReport(args.host, args.port, args.user, password)
    for receiver in args.to:
        report.add_receiver(receiver)

    try:
        report.set_email_content(
            args.subject,
            body,
            file_paths=[str(item.parent) for item in attachments],
            file_names=[item.name for item in attachments],
            is_html=args.html,
        )
    except (ValueError, FileNotFoundError) as exc:
        print(f"组装邮件失败: {exc}")
        return 1

    if args.dry_run:
        print(f"--dry-run: 未发送。主题: {args.subject}；收件人: {', '.join(args.to)}")
        return 0

    try:
        report.send_email()
    except MailError as exc:
        print(f"发送失败: {exc}")
        return 1
    print(f"已发送: {args.subject} -> {', '.join(args.to)}")
    return 0


def _cell(value: Any) -> str:
    if value is None:
        return "NULL"
    if isinstance(value, bytes):
        return value.hex()
    return str(value)


def _display_width(text: str) -> int:
    """按终端显示宽度计算（中日韩全角字符占两列）。"""
    return sum(2 if unicodedata.east_asian_width(char) in ("W", "F") else 1 for char in text)


def _print_rows(rows: Sequence[Any], output_format: str) -> None:
    if not rows:
        print("(0 行)")
        return

    first = rows[0]
    if isinstance(first, Mapping):
        header = [str(name) for name in first]
        grid = [[_cell(row.get(name)) for name in first] for row in rows]
    else:
        header = []
        grid = [[_cell(value) for value in row] for row in rows]

    if output_format == "json":
        payload = [dict(row) if isinstance(row, Mapping) else list(row) for row in rows]
        print(json.dumps(payload, ensure_ascii=False, indent=2, default=str))
        return

    if output_format == "csv":
        buffer = io.StringIO()
        writer = csv.writer(buffer, lineterminator="\n")
        writer.writerows([header] if header else [])
        writer.writerows(grid)
        print(buffer.getvalue(), end="")
        return

    table = [header] + grid if header else grid
    width = max(len(row) for row in table)
    padded = [list(row) + [""] * (width - len(row)) for row in table]
    widths = [max(_display_width(row[index]) for row in padded) for index in range(width)]
    lines = [
        "  ".join(row[index] + " " * (widths[index] - _display_width(row[index])) for index in range(width)).rstrip()
        for row in padded
    ]
    if header:
        lines.insert(1, "  ".join("-" * width_value for width_value in widths))
    print("\n".join(lines))


def _run_db(args: argparse.Namespace) -> int:
    if args.db_command != "query":
        return 1

    password = args.password or os.environ.get(_DB_PASSWORD_ENV)
    if not password:
        print(f"缺少密码: 用 --password 传入，或设置环境变量 {_DB_PASSWORD_ENV}。")
        return 1

    try:
        from wei_data_shu.database import MySQLDatabase, MySQLDatabaseError
    except ImportError:
        print(_DATABASE_HINT)
        return 1

    config = {
        "host": args.host,
        "port": args.port,
        "user": args.user,
        "password": password,
        "database": args.database,
    }
    try:
        with MySQLDatabase(config) as database:
            rows = database.fetch_query(args.sql, dictionary=True)
    except MySQLDatabaseError as exc:
        print(f"查询失败: {exc}")
        return 1

    _print_rows(rows, args.format)
    return 0


_COMMANDS = {
    "password": _run_password,
    "colors": _run_colors,
    "date": _run_date,
    "excel": _run_excel,
    "convert": _run_convert,
    "md": _run_md,
    "files": _run_files,
    "md2html": _run_md2html,
    "text": _run_text,
    "table": _run_table,
    "data": _run_data,
    "plot": _run_plot,
    "mail": _run_mail,
    "db": _run_db,
}


def main(argv: Sequence[str] | None = None) -> int:
    _ensure_utf8_stdout()
    parser = _build_parser()
    args = parser.parse_args(argv)

    handler = _COMMANDS.get(args.command)
    if handler is None:
        parser.print_help()
        return 0
    return handler(args)


__all__ = ["main"]
