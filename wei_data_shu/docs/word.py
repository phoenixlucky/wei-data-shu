"""Word (``.docx``) 读写，基于可选依赖 ``python-docx``。

安装::

    pip install wei-data-shu[docs]
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from ._deps import require_deps
from .model import (
    BulletList,
    CodeBlock,
    Document,
    Heading,
    Image,
    PageBreak,
    Paragraph,
    PathLike,
    Table,
)

_QUOTE_STYLE = "Quote"
_CODE_FONT = "Consolas"
_IMAGE_WIDTH_INCHES = 5.5


def _heading_level(style_name: str) -> int | None:
    """从段落样式名解析标题级别（兼容 ``Heading 1`` 与 ``标题 1``）。"""
    name = (style_name or "").strip()
    lowered = name.lower()
    if not (lowered.startswith("heading") or name.startswith("标题")):
        return None
    digits = re.findall(r"\d+", name)
    return int(digits[0]) if digits else 1


def _paragraph_text(paragraph: Any) -> str:
    return "".join(run.text for run in paragraph.runs).strip()


def write_docx(document: Document, path: PathLike) -> Path:
    """把 :class:`Document` 写入 ``.docx`` 文件，返回写入路径。"""
    require_deps("docx")
    from docx import Document as DocxDocument
    from docx.shared import Inches

    docx = DocxDocument()
    if document.title:
        docx.core_properties.title = document.title

    for node in document.nodes:
        if isinstance(node, Heading):
            docx.add_heading(node.text, level=node.level)
        elif isinstance(node, Paragraph):
            if node.style == "quote":
                docx.add_paragraph(node.text, style=_QUOTE_STYLE)
            else:
                docx.add_paragraph(node.text)
        elif isinstance(node, BulletList):
            style = "List Number" if node.ordered else "List Bullet"
            for item in node.items:
                docx.add_paragraph(item, style=style)
        elif isinstance(node, CodeBlock):
            paragraph = docx.add_paragraph()
            run = paragraph.add_run(node.code)
            run.font.name = _CODE_FONT
        elif isinstance(node, Table):
            rows = node.normalized()
            if rows:
                table = docx.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = "Table Grid"
                for row_index, row in enumerate(rows):
                    for col_index, value in enumerate(row):
                        table.cell(row_index, col_index).text = value
        elif isinstance(node, Image):
            if node.path and Path(node.path).exists():
                docx.add_picture(node.path, width=Inches(_IMAGE_WIDTH_INCHES))
                if node.caption:
                    docx.add_paragraph(node.caption, style="Caption")
        elif isinstance(node, PageBreak):
            docx.add_page_break()

    target = Path(path)
    docx.save(str(target))
    return target


def read_docx(path: PathLike) -> Document:
    """读取 ``.docx`` 文件为 :class:`Document`（按正文顺序保留标题/段落/列表/表格）。"""
    require_deps("docx")
    from docx import Document as DocxDocument
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph as DocxParagraph

    source = Path(path)
    docx = DocxDocument(str(source))
    document = Document(title=docx.core_properties.title or None)

    def append(node: Any) -> None:
        # 相邻同型列表合并，避免往返后被拆成碎片
        if (
            isinstance(node, BulletList)
            and document.nodes
            and isinstance(document.nodes[-1], BulletList)
            and document.nodes[-1].ordered == node.ordered
        ):
            document.nodes[-1].items.extend(node.items)
            return
        document.add(node)

    for child in docx.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = DocxParagraph(child, docx)
            text = _paragraph_text(paragraph)
            if not text:
                continue
            style_name = paragraph.style.name if paragraph.style is not None else ""
            level = _heading_level(style_name)
            if level is not None:
                append(Heading(text, level))
            elif style_name == _QUOTE_STYLE:
                append(Paragraph(text, "quote"))
            elif style_name.startswith("List Number"):
                append(BulletList([text], ordered=True))
            elif style_name.startswith("List Bullet"):
                append(BulletList([text]))
            else:
                append(Paragraph(text))
        elif tag == "tbl":
            table = DocxTable(child, docx)
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if rows:
                append(Table(rows, header=True))

    return document


__all__ = ["read_docx", "write_docx"]
