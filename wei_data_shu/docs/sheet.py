"""Excel (``.xlsx``) 读取：把工作簿转换为 :class:`~wei_data_shu.docs.model.Document`。

每个非空工作表渲染为一个二级标题（工作表名）+ 一个表格，因此
``convert("report.xlsx", "report.md")`` 可把整本工作簿搬进 Markdown / Word / PPT。

只读、不写：``.xlsx`` 在本包里仅作为转换的**源**格式。安装::

    pip install wei-data-shu[docs]

说明：以只读模式读取公式的缓存值；合并单元格只取左上角的值；全空行会被跳过。
"""

from __future__ import annotations

import datetime
from pathlib import Path
from typing import Any

from ._deps import require_deps
from .model import Document, Heading, PathLike, Table

SHEET_HEADING_LEVEL = 2

_MIDNIGHT = datetime.time(0, 0)


def _cell_text(value: Any) -> str:
    """把单元格值格式化为表格文本（日期去零时刻、整数值浮点去掉 ``.0``）。"""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime.datetime):
        if value.time() == _MIDNIGHT:
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, datetime.date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float):
        if value != value:  # NaN
            return ""
        if value.is_integer():
            return str(int(value))
    return str(value)


def _sheet_rows(worksheet: Any) -> list[list[str]]:
    """读取工作表为矩形文本行表；跳过全空行并裁掉尾部空行。"""
    rows: list[list[str]] = []
    for raw_row in worksheet.iter_rows(values_only=True):
        cells = [_cell_text(value) for value in raw_row]
        while cells and cells[-1] == "":
            cells.pop()
        if not cells:  # 全空行
            continue
        rows.append(cells)
    if not rows:
        return []
    width = max(len(row) for row in rows)
    return [row + [""] * (width - len(row)) for row in rows]


def read_xlsx(path: PathLike) -> Document:
    """读取 ``.xlsx`` 工作簿，每个非空工作表转为一个二级标题 + 表格。"""
    require_deps("openpyxl")
    from openpyxl import load_workbook

    source = Path(path)
    if not source.exists():
        raise FileNotFoundError(f"文件不存在: {source}")

    workbook = load_workbook(str(source), read_only=True, data_only=True)
    try:
        document = Document(title=source.stem)
        for worksheet in workbook.worksheets:
            rows = _sheet_rows(worksheet)
            if not rows:
                continue
            document.add(Heading(worksheet.title, SHEET_HEADING_LEVEL))
            document.add(Table(rows, header=True))
    finally:
        workbook.close()
    return document


__all__ = ["SHEET_HEADING_LEVEL", "read_xlsx"]
