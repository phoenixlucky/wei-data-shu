"""文档域示例：Markdown / Word / PowerPoint 生成、Excel 来源与跨格式互转。

运行前安装可选依赖（Word / PPT / Excel 来源部分需要）::

    pip install "wei-data-shu[docs]"

Markdown 与 HTML 预览零第三方依赖，只装核心包也可运行对应部分。
"""

from pathlib import Path

from wei_data_shu.docs import Document, Table, convert, df_to_markdown, parse_markdown, to_markdown

OUTPUT_DIR = Path("demo-output")

ROWS = [
    ["渠道", "销售额", "同比"],
    ["电商", 12580, "+12%"],
    ["门店", 9680, "+3%"],
    ["分销", 7320, "-5%"],
]


def build_report() -> Document:
    """用链式 API 组装一份月报。"""
    return (
        Document(title="月度销售报告")
        .add_heading("月度销售报告", 1)
        .add_paragraph("本报告由 wei_data_shu.docs 自动生成。")
        .add_heading("分渠道明细", 2)
        .add_table(ROWS)
        .add_bullets(["电商渠道同比 +12%，是增长主力", "分销渠道同比 -5%，需重点跟进"])
        .add_quote("数据口径：截至本月最后一个自然日。")
        .add_page_break()
        .add_heading("结语", 1)
        .add_paragraph("详见附件数据文件。")
    )


def build_sample_workbook(path: Path) -> Path:
    """造一个含两个工作表的工作簿，用于演示 xlsx 作为转换来源。"""
    from openpyxl import Workbook

    workbook = Workbook()
    summary = workbook.active
    summary.title = "汇总"
    for row in ROWS:
        summary.append(row)
    detail = workbook.create_sheet("明细")
    detail.append(["日期", "渠道", "金额"])
    detail.append(["2026-09-16", "电商", 12580])
    workbook.save(path)
    return path


def main() -> None:
    OUTPUT_DIR.mkdir(exist_ok=True)

    # 1) 一行把二维数据写成 Markdown
    print("已生成:", to_markdown(ROWS, OUTPUT_DIR / "quick.md", title="快速表格"))

    # 2) 链式构造后保存（按扩展名自动选择写入后端）
    report = build_report()
    md_path = report.save(OUTPUT_DIR / "report.md")
    print("已生成:", md_path)

    # 3) 跨格式互转：md -> docx / pptx（需要 wei-data-shu[docs]）
    for filename in ("report.docx", "report.pptx"):
        try:
            print("已生成:", convert(md_path, OUTPUT_DIR / filename))
        except ImportError as exc:
            print(f"跳过 {filename}: {exc}")

    # 4) 读回并取表格
    loaded = parse_markdown(md_path.read_text(encoding="utf-8"))
    table = next(node for node in loaded if isinstance(node, Table))
    print(df_to_markdown(table.rows))

    # 5) Excel 工作簿作为来源：整本读入（一个工作表一页），需要 openpyxl
    try:
        book = build_sample_workbook(OUTPUT_DIR / "sales.xlsx")
        print("已生成:", book)
        print("已生成:", convert(book, OUTPUT_DIR / "sales.md"))
    except ImportError as exc:
        print(f"跳过 xlsx 示例: {exc}")

    # 6) HTML 片段：notebook 中的 _repr_html_ 用的就是这个渲染器（零依赖）
    print("HTML 预览（截断）:", report.to_html()[:140], "...")


if __name__ == "__main__":
    main()
