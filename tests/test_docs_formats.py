"""文档域测试：Markdown / HTML 渲染、Excel 源、跨格式互转、一行式 API 与 Jupyter 兼容。"""

import io
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from wei_data_shu.cli import main as cli_main
from wei_data_shu.docs import (
    BulletList,
    CodeBlock,
    Document,
    Heading,
    PageBreak,
    Paragraph,
    Table,
    build,
    convert,
    df_to_markdown,
    markdown_to_df,
    parse_markdown,
    read_doc,
    read_markdown,
    read_xlsx,
    render_html,
    render_markdown,
    to_markdown,
    write_markdown,
)
from wei_data_shu.docs._deps import has_deps, require_deps
from wei_data_shu.utils import in_notebook

SAMPLE_MD = """# 月度报告

本月销售额同比增长 12%，
环比持平。

## 分渠道

- 电商 12580
- 门店 9680

| 渠道 | 销售额 |
| --- | --- |
| 电商 | 12580 |
| 门店 | 9680 |

> 备注：数据截至月末。

```python
print("hello")
```

---

完
"""


def _write_sample_workbook(path):
    """生成含 有数据 / 全空 / 有数据 三个工作表、带日期与空行的样例工作簿。"""
    import datetime

    from openpyxl import Workbook

    workbook = Workbook()
    daily = workbook.active
    daily.title = "日报"
    daily.append(["日期", "渠道", "销售额"])
    daily.append([datetime.datetime(2026, 9, 16), "电商", 12580.0])
    daily.append([datetime.datetime(2026, 9, 16, 8, 30), "门店", 9680.5])
    daily.append([None, None, None])
    workbook.create_sheet("空表")
    notes = workbook.create_sheet("备注")
    notes.append(["说明"])
    notes.append(["仅一行"])
    workbook.save(path)


class TestMarkdownParsing(unittest.TestCase):
    def test_parse_node_types_in_order(self):
        document = parse_markdown(SAMPLE_MD)
        kinds = [type(node).__name__ for node in document.nodes]
        self.assertEqual(
            kinds,
            [
                "Heading",
                "Paragraph",
                "Heading",
                "BulletList",
                "Table",
                "Paragraph",
                "CodeBlock",
                "PageBreak",
                "Paragraph",
            ],
        )

    def test_heading_level_and_table_content(self):
        document = parse_markdown(SAMPLE_MD)
        self.assertEqual(document[0], Heading("月度报告", 1))
        table = next(node for node in document if isinstance(node, Table))
        self.assertEqual(table.rows[0], ["渠道", "销售额"])
        self.assertEqual(table.rows[2], ["门店", "9680"])

    def test_quote_and_code_block(self):
        document = parse_markdown(SAMPLE_MD)
        quote = next(node for node in document if isinstance(node, Paragraph) and node.style == "quote")
        self.assertIn("数据截至月末", quote.text)
        code = next(node for node in document if isinstance(node, CodeBlock))
        self.assertEqual(code.language, "python")
        self.assertEqual(code.code, 'print("hello")')

    def test_cjk_lines_joined_without_space(self):
        document = parse_markdown("本月销售额同比增长 12%，\n环比持平。\n")
        self.assertEqual(document[0].text, "本月销售额同比增长 12%，环比持平。")

    def test_render_round_trip_is_stable(self):
        once = render_markdown(parse_markdown(SAMPLE_MD))
        twice = render_markdown(parse_markdown(once))
        self.assertEqual(once, twice)
        self.assertIn("| 渠道 | 销售额 |", once)
        self.assertIn("---", once)


class TestMarkdownFiles(unittest.TestCase):
    def test_write_and_read_markdown(self):
        with TemporaryDirectory() as tmp:
            target = Path(tmp) / "report.md"
            document = (
                Document(title="月报")
                .add_heading("月报", 1)
                .add_paragraph("概览")
                .add_bullets(["电商", "门店"])
                .add_table([["渠道", "销售额"], ["电商", "12580"]])
            )
            write_markdown(document, target)
            self.assertTrue(target.exists())

            loaded = read_markdown(target)
            self.assertEqual(loaded[0], Heading("月报", 1))
            self.assertEqual(loaded[-1].rows[1], ["电商", "12580"])

    def test_document_save_and_open_dispatch(self):
        with TemporaryDirectory() as tmp:
            target = Path(tmp) / "auto.md"
            Document().add_heading("标题", 2).save(target)
            self.assertEqual(Document.open(target)[0], Heading("标题", 2))

    def test_unsupported_suffix_raises(self):
        with TemporaryDirectory() as tmp:
            with self.assertRaises(ValueError):
                Document().save(Path(tmp) / "report.xyz")


class TestTableHelpers(unittest.TestCase):
    def test_df_to_markdown_renders_header_separator(self):
        text = df_to_markdown([["渠道", "销售额"], ["电商", 12580]])
        self.assertEqual(
            text.splitlines(),
            ["| 渠道 | 销售额 |", "| --- | --- |", "| 电商 | 12580 |"],
        )

    def test_df_to_markdown_accepts_dataframe_like(self):
        class FakeFrame:
            columns = ["a", "b"]

            def itertuples(self, index=False, name=None):
                return [(1, 2), (3, 4)]

        text = df_to_markdown(FakeFrame())
        self.assertIn("| a | b |", text)
        self.assertIn("| 3 | 4 |", text)

    def test_markdown_to_df_records_without_pandas(self):
        records = markdown_to_df(SAMPLE_MD, as_records=True)
        self.assertEqual(records[0], {"渠道": "电商", "销售额": "12580"})
        self.assertEqual(len(records), 2)

    def test_markdown_to_df_without_table_raises(self):
        with self.assertRaises(ValueError):
            markdown_to_df("# 只有标题\n")


class TestQuickApi(unittest.TestCase):
    def test_build_from_rows_and_dict(self):
        document = build([["渠道", "销售额"], ["电商", "12580"]], title="月报")
        self.assertEqual(document.title, "月报")
        self.assertIsInstance(document[0], Heading)
        self.assertIsInstance(document[1], Table)

        nested = build({"分渠道": ["电商", "门店"]})
        self.assertEqual(nested[0], Heading("分渠道", 2))
        self.assertEqual(nested[1], BulletList(["电商", "门店"]))

    def test_to_markdown_returns_text_or_path(self):
        text = to_markdown([["a", "b"], ["1", "2"]])
        self.assertIn("| a | b |", text)
        with TemporaryDirectory() as tmp:
            target = Path(tmp) / "auto.md"
            self.assertEqual(to_markdown([["a"], ["1"]], target), target)
            self.assertTrue(target.exists())

    def test_document_markdown_repr_for_jupyter(self):
        document = Document().add_heading("标题", 1)
        self.assertEqual(document._repr_markdown_(), document.to_markdown())
        self.assertEqual(repr(document), "Document(title=None, nodes=1)")


class TestHtmlRendering(unittest.TestCase):
    def test_table_renders_as_html_table(self):
        html = render_html(Document().add_table([["渠道", "销售额"], ["电商", "12580"]]))
        self.assertIn("<table", html)
        self.assertIn("<th", html)
        self.assertIn("<td", html)
        self.assertIn("电商", html)

    def test_text_is_escaped(self):
        document = (
            Document()
            .add_paragraph("<script>alert(1)</script>")
            .add_table([["<b>表头</b>"], ["<i>值</i>"]])
        )
        html = render_html(document)
        self.assertNotIn("<script>", html)
        self.assertIn("&lt;script&gt;", html)
        self.assertNotIn("<b>表头</b>", html)
        self.assertIn("&lt;b&gt;表头&lt;/b&gt;", html)

    def test_node_tags(self):
        document = (
            Document()
            .add_heading("标题", 2)
            .add_bullets(["a", "b"])
            .add_bullets(["1"], ordered=True)
            .add_quote("引用")
            .add_code('print("hi")', "python")
            .add_page_break()
        )
        html = render_html(document)
        self.assertIn("<h2>标题</h2>", html)
        self.assertIn("<ul><li>a</li><li>b</li></ul>", html)
        self.assertIn("<ol><li>1</li></ol>", html)
        self.assertIn("<blockquote>", html)
        self.assertIn('<code class="language-python">', html)
        self.assertIn("<hr>", html)

    def test_paragraph_newline_becomes_br(self):
        html = render_html(Document().add_paragraph("第一行\n第二行"))
        self.assertIn("第一行<br>第二行", html)

    def test_image_attribute_is_escaped(self):
        html = render_html(Document().add_image('chart".png', caption="图"))
        self.assertIn("chart&quot;.png", html)
        self.assertNotIn('src="chart".png"', html)

    def test_repr_html_and_title_option(self):
        document = Document(title="月报").add_heading("正文", 1)
        self.assertEqual(document._repr_html_(), document.to_html())
        self.assertNotIn("月报", document.to_html())
        self.assertIn("月报", document.to_html(include_title=True))


class TestOptionalDependencyGuard(unittest.TestCase):
    def test_missing_dependency_message_points_to_extra(self):
        with patch("wei_data_shu.docs._deps._is_available", return_value=False):
            with self.assertRaises(ImportError) as ctx:
                require_deps("docx", "pptx")
            message = str(ctx.exception)
            self.assertIn("wei-data-shu[docs]", message)
            self.assertIn("docx", message)

    def test_missing_openpyxl_message_points_to_extra(self):
        with patch("wei_data_shu.docs._deps._is_available", return_value=False):
            with self.assertRaises(ImportError) as ctx:
                require_deps("openpyxl")
            self.assertIn("wei-data-shu[docs]", str(ctx.exception))

    def test_has_deps_reports_installation_state(self):
        self.assertEqual(has_deps("docx"), has_deps("docx"))

    def test_dependency_probe_does_not_import_optional_modules(self):
        """惰性契约：可用性探测不得在模块级 import 可选依赖。"""
        from wei_data_shu.docs import _deps

        for name in ("docx", "pptx", "openpyxl"):
            self.assertFalse(hasattr(_deps, name))


class TestExcelSource(unittest.TestCase):
    def test_xlsx_cannot_be_a_target(self):
        with TemporaryDirectory() as tmp:
            source = Path(tmp) / "report.md"
            source.write_text("# 标题\n", encoding="utf-8")
            with self.assertRaises(ValueError) as ctx:
                convert(source, Path(tmp) / "out.xlsx")
            self.assertIn(".xlsx", str(ctx.exception))

    def test_unknown_source_suffix_raises(self):
        with TemporaryDirectory() as tmp:
            source = Path(tmp) / "data.xyz"
            source.write_text("x", encoding="utf-8")
            with self.assertRaises(ValueError):
                convert(source, Path(tmp) / "out.md")

    @unittest.skipUnless(has_deps("openpyxl"), "需要 wei-data-shu[docs]")
    def test_read_xlsx_multi_sheet_and_formatting(self):
        with TemporaryDirectory() as tmp:
            book = Path(tmp) / "sales.xlsx"
            _write_sample_workbook(book)

            document = read_xlsx(book)
            self.assertEqual(document.title, "sales")
            # “空表” 被跳过，只保留两个有内容的工作表
            self.assertEqual(
                [type(node).__name__ for node in document],
                ["Heading", "Table", "Heading", "Table"],
            )
            self.assertEqual(document[0], Heading("日报", 2))
            self.assertEqual(document[2], Heading("备注", 2))
            # 全空行被跳过；日期去掉零时刻；整数值浮点去掉 .0；小数保留
            self.assertEqual(
                document[1].rows,
                [
                    ["日期", "渠道", "销售额"],
                    ["2026-09-16", "电商", "12580"],
                    ["2026-09-16 08:30:00", "门店", "9680.5"],
                ],
            )

    @unittest.skipUnless(has_deps("openpyxl"), "需要 wei-data-shu[docs]")
    def test_convert_xlsx_to_markdown(self):
        with TemporaryDirectory() as tmp:
            book = Path(tmp) / "sales.xlsx"
            _write_sample_workbook(book)

            text = convert(book, Path(tmp) / "sales.md").read_text(encoding="utf-8")
            self.assertIn("## 日报", text)
            self.assertIn("## 备注", text)
            self.assertNotIn("空表", text)
            self.assertIn("| 日期 | 渠道 | 销售额 |", text)
            self.assertIn("| 2026-09-16 | 电商 | 12580 |", text)

    @unittest.skipUnless(has_deps("openpyxl"), "需要 wei-data-shu[docs]")
    def test_read_doc_dispatches_xlsx(self):
        with TemporaryDirectory() as tmp:
            book = Path(tmp) / "sales.xlsx"
            _write_sample_workbook(book)
            self.assertEqual(read_doc(book)[0], Heading("日报", 2))

    @unittest.skipUnless(has_deps("openpyxl"), "需要 wei-data-shu[docs]")
    def test_cli_md_reads_xlsx(self):
        with TemporaryDirectory() as tmp:
            book = Path(tmp) / "sales.xlsx"
            _write_sample_workbook(book)
            buffer = io.StringIO()
            with redirect_stdout(buffer):
                code = cli_main(["md", str(book)])
            self.assertEqual(code, 0)
            self.assertIn("## 日报", buffer.getvalue())


class TestWordRoundTrip(unittest.TestCase):
    @unittest.skipUnless(has_deps("docx"), "需要 wei-data-shu[docs]")
    def test_markdown_to_docx_and_back(self):
        with TemporaryDirectory() as tmp:
            source = Path(tmp) / "report.md"
            target = Path(tmp) / "report.docx"
            write_markdown(parse_markdown(SAMPLE_MD), source)
            convert(source, target)
            self.assertTrue(target.exists())

            text = convert(target, Path(tmp) / "back.md").read_text(encoding="utf-8")
            self.assertIn("# 月度报告", text)
            self.assertIn("- 电商 12580", text)
            self.assertIn("| 渠道 | 销售额 |", text)


class TestSlidesRoundTrip(unittest.TestCase):
    @unittest.skipUnless(has_deps("pptx"), "需要 wei-data-shu[docs]")
    def test_markdown_to_pptx_and_back(self):
        with TemporaryDirectory() as tmp:
            source = Path(tmp) / "report.md"
            target = Path(tmp) / "report.pptx"
            write_markdown(parse_markdown(SAMPLE_MD), source)
            convert(source, target)
            self.assertTrue(target.exists())

            document = read_doc(target)
            headings = [node.text for node in document if isinstance(node, Heading)]
            self.assertIn("月度报告", headings)
            self.assertIn("分渠道", headings)


class TestNotebookCompat(unittest.TestCase):
    def test_in_notebook_false_in_plain_interpreter(self):
        self.assertFalse(in_notebook())

    def test_read_doc_dispatches_by_extension(self):
        with TemporaryDirectory() as tmp:
            target = Path(tmp) / "note.md"
            target.write_text("# 标题\n", encoding="utf-8")
            self.assertEqual(read_doc(target)[0], Heading("标题", 1))


class TestCliDocumentCommands(unittest.TestCase):
    def test_md_command_prints_markdown(self):
        with TemporaryDirectory() as tmp:
            source = Path(tmp) / "report.md"
            source.write_text("# 标题\n\n正文\n", encoding="utf-8")
            buffer = io.StringIO()
            with redirect_stdout(buffer):
                code = cli_main(["md", str(source)])
            self.assertEqual(code, 0)
            self.assertIn("# 标题", buffer.getvalue())

    def test_convert_command_writes_target(self):
        with TemporaryDirectory() as tmp:
            source = Path(tmp) / "report.md"
            target = Path(tmp) / "copy.md"
            source.write_text("# 标题\n", encoding="utf-8")
            buffer = io.StringIO()
            with redirect_stdout(buffer):
                code = cli_main(["convert", str(source), "-o", str(target)])
            self.assertEqual(code, 0)
            self.assertIn("已生成", buffer.getvalue())
            self.assertEqual(target.read_text(encoding="utf-8").strip(), "# 标题")

    def test_convert_command_reports_failure(self):
        buffer = io.StringIO()
        with redirect_stdout(buffer):
            code = cli_main(["convert", "missing.md", "-o", "out.md"])
        self.assertEqual(code, 1)
        self.assertIn("转换失败", buffer.getvalue())


if __name__ == "__main__":
    unittest.main()
